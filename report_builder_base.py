"""
Seminar Report Document Generator
=================================
Topic: Comparative Analysis of Reinforcement Learning Strategies (Proximal Policy Optimization vs Deep Q-Networks)
       for Adaptive AI Agent Orchestration & Sequential Decision Making
Author: Aniska Bachar (Roll No. 05)
Institution: Thakur College of Engineering & Technology, Mumbai
Department: Department of Artificial Intelligence & Machine Learning
Degree: Bachelor of Technology (NOT ENGINEERING)
Academic Year: 2025-26
Guide: Dr. Shiwani Gupta (HOD), Professor & Head of Department
Co-Guide: Mrs. Seema Nagrani, Assistant Professor
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set cell padding in twips (1 pt = 20 twips)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}>'
                      f'<w:top w:w="{top}" w:type="dxa"/>'
                      f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
                      f'<w:left w:w="{left}" w:type="dxa"/>'
                      f'<w:right w:w="{right}" w:type="dxa"/>'
                      f'</w:tcMar>')
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    """Apply background fill color to cell"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    """Set elegant, clean table borders (booktabs style)"""
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}>'
                        f'<w:top w:val="{val}" w:sz="8" w:space="0" w:color="1B365D"/>'
                        f'<w:bottom w:val="{val}" w:sz="8" w:space="0" w:color="1B365D"/>'
                        f'<w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:insideV w:val="none"/>'
                        f'<w:left w:val="none"/>'
                        f'<w:right w:val="none"/>'
                        f'</w:tblBorders>')
    tblPr.append(borders)

def add_page_number_to_footer(footer):
    """Add dynamic page numbering to footer"""
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run("Page ")
    run.font.name = "Times New Roman"
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = "PAGE"
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    p._p.append(fld1)
    p._p.append(instr)
    p._p.append(fld2)
    p._p.append(fld3)

    run2 = p.add_run(" of ")
    run2.font.name = "Times New Roman"
    run2.font.size = Pt(10)
    run2.font.color.rgb = RGBColor(0x71, 0x80, 0x96)

    fld4 = OxmlElement('w:fldChar')
    fld4.set(qn('w:fldCharType'), 'begin')
    instr2 = OxmlElement('w:instrText')
    instr2.set(qn('xml:space'), 'preserve')
    instr2.text = "NUMPAGES"
    fld5 = OxmlElement('w:fldChar')
    fld5.set(qn('w:fldCharType'), 'separate')
    fld6 = OxmlElement('w:fldChar')
    fld6.set(qn('w:fldCharType'), 'end')
    p._p.append(fld4)
    p._p.append(instr2)
    p._p.append(fld5)
    p._p.append(fld6)

print("Base helper definitions loaded successfully.")
